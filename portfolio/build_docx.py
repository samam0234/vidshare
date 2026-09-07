"""VidShare-포트폴리오.md -> VidShare-포트폴리오.docx

    pip install python-docx
    python portfolio/build_docx.py

마크다운 전체를 지원하는 변환기가 아니라, 이 문서가 실제로 쓰는 문법
(제목 / 표 / 코드블록 / 목록 / 인용 / 굵게 / 인라인코드 / 구분선)만 다룬다.
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches

ROOT = Path(__file__).resolve().parent
SRC = ROOT / "VidShare-포트폴리오.md"
OUT = ROOT / "VidShare-포트폴리오.docx"

BODY_FONT = "맑은 고딕"
CODE_FONT = "Consolas"

INK = RGBColor(0x1A, 0x1D, 0x24)
MUTED = RGBColor(0x5A, 0x63, 0x72)
ACCENT = RGBColor(0x5B, 0x3D, 0xF5)
CODE_INK = RGBColor(0x22, 0x33, 0x55)
RULE = "D6DAE2"
CODE_BG = "F4F5F8"
HEAD_BG = "EEF0F6"


def set_cell_bg(cell, hex_color: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(shd)


def shade_paragraph(par, hex_color: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    par._p.get_or_add_pPr().append(shd)


def set_east_asian(run) -> None:
    """python-docx 는 동아시아 글꼴을 따로 지정해야 한글에 실제로 적용된다."""
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), run.font.name or BODY_FONT)


def add_run(par, text: str, *, bold=False, code=False, color=None, size=10.5):
    run = par.add_run(text)
    run.bold = bold
    run.font.name = CODE_FONT if code else BODY_FONT
    run.font.size = Pt(size - (0.5 if code else 0))
    run.font.color.rgb = color or (CODE_INK if code else INK)
    set_east_asian(run)
    return run


INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")


def add_inline(par, text: str, *, bold=False, color=None, size=10.5):
    """**굵게** 와 `인라인 코드` 만 해석한다."""
    for piece in INLINE.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**"):
            add_run(par, piece[2:-2], bold=True, color=color, size=size)
        elif piece.startswith("`") and piece.endswith("`"):
            add_run(par, piece[1:-1], code=True, size=size)
        else:
            add_run(par, piece, bold=bold, color=color, size=size)


def horizontal_rule(doc) -> None:
    par = doc.add_paragraph()
    par.paragraph_format.space_before = Pt(6)
    par.paragraph_format.space_after = Pt(10)
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:color"), RULE)
    pbdr.append(bottom)
    par._p.get_or_add_pPr().append(pbdr)


def heading(doc, text: str, level: int) -> None:
    sizes = {1: 20, 2: 15, 3: 12.5, 4: 11}
    par = doc.add_paragraph()
    par.paragraph_format.space_before = Pt({1: 18, 2: 16, 3: 12, 4: 10}[level])
    par.paragraph_format.space_after = Pt(6)
    par.paragraph_format.keep_with_next = True
    color = ACCENT if level <= 2 else INK
    add_inline(par, text, bold=True, color=color, size=sizes[level])


def code_block(doc, lines: list[str]) -> None:
    for line in lines or [""]:
        par = doc.add_paragraph()
        par.paragraph_format.space_before = Pt(0)
        par.paragraph_format.space_after = Pt(0)
        par.paragraph_format.left_indent = Inches(0.18)
        shade_paragraph(par, CODE_BG)
        add_run(par, line if line else " ", code=True)


def table_block(doc, rows: list[list[str]]) -> None:
    header, body = rows[0], rows[1:]
    table = doc.add_table(rows=len(rows), cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for col, text in enumerate(header):
        cell = table.cell(0, col)
        cell.text = ""
        set_cell_bg(cell, HEAD_BG)
        par = cell.paragraphs[0]
        par.paragraph_format.space_before = Pt(3)
        par.paragraph_format.space_after = Pt(3)
        add_inline(par, text, bold=True, size=9.5)

    for r, row in enumerate(body, start=1):
        for col in range(len(header)):
            cell = table.cell(r, col)
            cell.text = ""
            par = cell.paragraphs[0]
            par.paragraph_format.space_before = Pt(2)
            par.paragraph_format.space_after = Pt(2)
            add_inline(par, row[col] if col < len(row) else "", size=9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)


BLOCK_START = re.compile(r"^(#{1,4}\s|>\s|\||```|-{3,}$|\s*[-*]\s|\s*\d+\.\s)")


def starts_block(line: str) -> bool:
    return not line.strip() or bool(BLOCK_START.match(line))


def gather(lines: list[str], i: int) -> tuple[str, int]:
    """마크다운에서 줄바꿈 하나는 문단을 끊지 않는다.
    빈 줄이나 새 블록이 나올 때까지 이어 붙인다."""
    buf = [lines[i].strip()]
    i += 1
    while i < len(lines) and not starts_block(lines[i]):
        buf.append(lines[i].strip())
        i += 1
    return " ".join(buf), i


def split_row(line: str) -> list[str]:
    # 셀 안의 \| 는 구분자가 아니라 문자다
    cells = re.split(r"(?<!\\)\|", line.strip().strip("|"))
    return [c.strip().replace("\\|", "|") for c in cells]


def is_divider(line: str) -> bool:
    return bool(re.fullmatch(r"\|[\s:|-]+\|", line.strip()))


def cover(doc) -> None:
    for _ in range(4):
        doc.add_paragraph()

    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(par, "VidShare", bold=True, color=ACCENT, size=34)

    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(par, "영상 공유 플랫폼 · 풀스택 개인 프로젝트", color=MUTED, size=13)

    doc.add_paragraph()
    horizontal_rule(doc)

    for label, value in [
        ("구성", "사용자 웹앱 · 관리자 콘솔 · REST API 서버"),
        ("기술", "Next.js 16 · React 19 · TypeScript · Express · SQLite · LangChain"),
        ("기간", "2026-08-14 ~ 2026-09-07"),
        ("규모", "약 15,700줄 · API 82개 · 테이블 22개 · 테스트 177건"),
    ]:
        par = doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.space_after = Pt(3)
        add_run(par, f"{label}   ", bold=True, color=MUTED, size=10)
        add_run(par, value, color=INK, size=10)

    doc.add_page_break()


def convert(md: str) -> Document:
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = BODY_FONT
    style.font.size = Pt(10.5)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.35

    for section in doc.sections:
        section.left_margin = section.right_margin = Inches(0.9)
        section.top_margin = section.bottom_margin = Inches(0.9)

    cover(doc)

    lines = md.splitlines()
    i = 0
    seen_title = False

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if stripped.startswith("```"):
            i += 1
            buf = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1
            code_block(doc, buf)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue

        if stripped.startswith("|") and i + 1 < len(lines) and is_divider(lines[i + 1]):
            rows = [split_row(stripped)]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(split_row(lines[i]))
                i += 1
            table_block(doc, rows)
            continue

        if re.fullmatch(r"-{3,}", stripped):
            horizontal_rule(doc)
            i += 1
            continue

        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            # 문서 첫 제목은 표지가 대신한다
            if level == 1 and not seen_title:
                seen_title = True
                i += 1
                continue
            heading(doc, m.group(2), level)
            i += 1
            continue

        if stripped.startswith("> "):
            buf = []
            while i < len(lines) and lines[i].strip().startswith("> "):
                buf.append(lines[i].strip()[2:])
                i += 1
            par = doc.add_paragraph()
            par.paragraph_format.left_indent = Inches(0.25)
            add_inline(par, " ".join(buf), color=MUTED)
            continue

        m = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)$", line)
        if m:
            bullet = m.group(2) in ("-", "*")
            depth = len(m.group(1)) // (2 if bullet else 3)
            # 이어지는 들여쓴 줄은 같은 항목의 연속이다
            buf = [m.group(3).strip()]
            i += 1
            while i < len(lines) and lines[i].strip() and not starts_block(lines[i]):
                buf.append(lines[i].strip())
                i += 1
            par = doc.add_paragraph(style="List Bullet" if bullet else "List Number")
            par.paragraph_format.left_indent = Inches(0.25 + 0.25 * depth)
            par.paragraph_format.space_after = Pt(2)
            add_inline(par, " ".join(buf))
            continue

        text, i = gather(lines, i)
        par = doc.add_paragraph()
        add_inline(par, text)

    return doc


def main() -> int:
    if not SRC.exists():
        print(f"원본을 찾을 수 없다: {SRC}", file=sys.stderr)
        return 1
    doc = convert(SRC.read_text(encoding="utf-8"))
    doc.save(OUT)
    print(f"생성: {OUT}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
